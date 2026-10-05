"""Extrae el cuestionario COSO ERM del libro de Excel original (SECI) a JSON.

Uso:
    python coso_erm/herramientas/extraer_excel.py RUTA/1._Eval_COSO_ERM_2026.xlsx

Genera ``coso_erm/datos/fuente_excel.json`` con:
  * componentes -> principios -> puntos de interés (preguntas) tal como vienen
    en las hojas GC, EST, DES, REV e ICR;
  * glosario, tabla de madurez y normativa;
  * respuestas de ejemplo (atributos a-h, evidencias, plan de solución), que el
    sistema ofrece como evaluación de demostración.

El archivo resultante es la entrada de ``construir.py``, que aplica las
actualizaciones 2026 y arma el HTML autocontenido.
"""
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci, get_column_letter

COMPONENTES = [
    ("GC", "Gobierno y Cultura"),
    ("EST", "Estrategia y Establecimiento de Objetivos"),
    ("DES", "Desempeño"),
    ("REV", "Revisión y Ajuste"),
    ("ICR", "Información, Comunicación y Reporte"),
]
SALIDA = Path(__file__).resolve().parent.parent / "datos" / "fuente_excel.json"


def valor(ws, fila, col):
    x = ws.cell(fila, ci(col)).value
    if x is None:
        return ""
    if isinstance(x, (datetime.date, datetime.datetime)):
        return x.strftime("%Y-%m-%d")
    s = str(x).strip()
    return "" if s == "$" else s


def extraer(ruta):
    wb = openpyxl.load_workbook(ruta)
    componentes, demo = [], {}
    for codigo, nombre in COMPONENTES:
        ws = wb[codigo]
        # La hoja GC tiene una columna menos que las demás antes de los atributos.
        off = 0 if codigo == "GC" else 1
        col = lambda c: get_column_letter(ci(c) + off)  # noqa: E731
        principios, actual = [], None
        for r in range(5, ws.max_row + 1):
            b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
            if isinstance(b, str) and b.startswith("Principio"):
                m = re.match(r"Principio\s+(\d+):\s*(.*)", b)
                actual = {"n": int(m.group(1)), "t": m.group(2).strip(), "q": []}
                principios.append(actual)
            elif isinstance(b, str) and re.match(r"^\d+\.\d+\.\d+", b):
                qid = b.strip()
                texto = re.sub(r"^\d+\.\s*", "", (c or "").strip())
                actual["q"].append({"id": qid, "t": texto})
                attrs = [ws.cell(r, ci(col("J")) + k).value for k in range(8)]
                attrs = [int(a) if a in (0, 1) else None for a in attrs]
                ev_cols = ["V", "W", "X", "Y", "Z", "AA", "AB"] if off == 0 else ["W", "X", "Y", "Z", "AA", "AB", "AC"]
                d = {
                    "a": attrs,
                    "area": valor(ws, r, col("H")),
                    "cargo": valor(ws, r, col("I")),
                    "ev": [valor(ws, r, x) for x in ev_cols],
                    "ubic": valor(ws, r, "AC" if off == 0 else "AD"),
                    "matriz": valor(ws, r, "AD" if off == 0 else "AE"),
                    "prio": valor(ws, r, "AE") if off == 0 else "",
                    "coment": valor(ws, r, "AF"),
                    "defi": valor(ws, r, "AH"),
                    "plan": valor(ws, r, "AI"),
                    "resp": valor(ws, r, "AJ"),
                    "fi": valor(ws, r, "AK"),
                    "ff": valor(ws, r, "AL"),
                    "presup": valor(ws, r, "AM"),
                    "entreg": valor(ws, r, "AN"),
                }
                d = {k: v for k, v in d.items()
                     if v and v != [""] * 7 and not (k == "a" and all(y is None for y in v))}
                if d:
                    demo[qid] = d
        principios.sort(key=lambda p: p["n"])
        componentes.append({"code": codigo, "name": nombre, "principles": principios})

    def pares(hoja, ck, cv, ini, fin):
        h = wb[hoja]
        return [[str(h[f"{ck}{r}"].value).strip(), str(h[f"{cv}{r}"].value or "").strip()]
                for r in range(ini, fin + 1) if h[f"{ck}{r}"].value]

    mad = wb["Madurez"]
    return {
        "components": componentes,
        "glosario": pares("Glosario", "B", "H", 9, 40),
        "madurez": [[mad[f"B{r}"].value, str(mad[f"F{r}"].value), mad[f"H{r}"].value] for r in range(10, 21, 2)],
        "normativa": [x[1] for x in pares("Normativa", "B", "D", 8, 37)],
        "demo": demo,
    }


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    datos = extraer(sys.argv[1])
    SALIDA.write_text(json.dumps(datos, ensure_ascii=False, indent=1), encoding="utf-8")
    n = sum(len(p["q"]) for c in datos["components"] for p in c["principles"])
    print(f"{n} puntos de interés extraídos -> {SALIDA}")
